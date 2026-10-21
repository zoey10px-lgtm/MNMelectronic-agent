"""Extract text and embedded images from customer PDF/DOCX uploads.

No network calls or LLMs. Extracted images can be sent separately to local YOLO.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

MAX_BYTES = 12 * 1024 * 1024
MAX_PAGES = 40
MAX_IMAGES = 12
MAX_TEXT = 30000


def preprocess_document(filename: str, data: bytes) -> dict:
    """Return text and images as bytes; callers must not JSON-serialize raw images."""
    if len(data) > MAX_BYTES:
        raise ValueError("Document exceeds 12 MB.")
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        import fitz
        try:
            doc = fitz.open(stream=data, filetype="pdf")
            if doc.needs_pass:
                raise ValueError("Password-protected PDF is not supported.")
            if len(doc) > MAX_PAGES:
                raise ValueError("PDF exceeds 40 pages.")
            chunks, images = [], []
            for page in doc:
                chunks.append(page.get_text("text"))
                for entry in page.get_images(full=True):
                    if len(images) >= MAX_IMAGES:
                        break
                    raw = doc.extract_image(entry[0])["image"]
                    _append_image(images, raw)
            doc.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Unable to parse PDF.") from exc
    elif ext == ".docx":
        from docx import Document
        try:
            doc = Document(io.BytesIO(data))
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text for cell in row.cells))
            images = []
            for rel in doc.part.rels.values():
                if len(images) >= MAX_IMAGES:
                    break
                if "image" in rel.reltype:
                    _append_image(images, rel.target_part.blob)
        except Exception as exc:
            raise ValueError("Unable to parse DOCX.") from exc
    else:
        raise ValueError("Only PDF and DOCX are supported.")
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(chunks)).strip()[:MAX_TEXT]
    return {"filename": Path(filename).name, "type": ext[1:],
            "text": text, "images": images, "image_count": len(images)}


def _append_image(images: list, raw: bytes) -> None:
    from PIL import Image, ImageOps, UnidentifiedImageError
    if len(raw) > 8 * 1024 * 1024:
        return
    try:
        with Image.open(io.BytesIO(raw)) as source:
            if source.width * source.height > 20_000_000:
                return
            image = ImageOps.exif_transpose(source).convert("RGB")
            output = io.BytesIO()
            image.save(output, format="JPEG", quality=90)
            images.append(output.getvalue())
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
        return
